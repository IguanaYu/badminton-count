from io import BytesIO
from typing import List

from openpyxl import Workbook
from sqlalchemy.orm import Session

from ..models import Match, MatchParticipant, User
from ..stats.services import compute_player_summary


def _workbook_to_bytes(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()


def export_player_matches(db: Session, user: User) -> bytes:
    matches: List[Match] = (
        db.query(Match)
        .join(MatchParticipant)
        .filter(MatchParticipant.user_id == user.id)
        .order_by(Match.played_at.asc())
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Matches"
    ws.append([
        "Match ID",
        "Played At",
        "Session",
        "Type",
        "Side",
        "Score A",
        "Score B",
        "Winner",
        "Teammates",
        "Opponents",
        "Recorder",
    ])

    for match in matches:
        user_participant = next(p for p in match.participants if p.user_id == user.id)
        user_side = user_participant.side
        teammates = ", ".join(
            p.user.full_name for p in match.participants if p.side == user_side and p.user_id != user.id
        )
        opponents = ", ".join(p.user.full_name for p in match.participants if p.side != user_side)
        ws.append(
            [
                match.id,
                match.played_at.isoformat(sep=" "),
                match.session_number or "-",
                match.match_type.value,
                user_side,
                match.score_a,
                match.score_b,
                match.winner_side,
                teammates or "-",
                opponents,
                match.recorded_by.full_name,
            ]
        )

    return _workbook_to_bytes(wb)


def export_player_summary(db: Session, user: User) -> bytes:
    summary = compute_player_summary(db, user)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    ws.append(["Player", user.full_name])
    ws.append(["Username", user.username])
    ws.append(["Gender", user.gender.value])
    ws.append([])
    ws.append(["Total Matches", summary["total_matches"]])
    ws.append(["Wins", summary["wins"]])
    ws.append(["Losses", summary["losses"]])
    ws.append(["Win Rate", summary["win_rate"]])

    ws.append([])
    ws.append(["Category", "Matches", "Wins", "Losses", "Win Rate"])
    for key in ["singles", "doubles", "one_vs_two_as_one", "one_vs_two_as_two"]:
        stats = summary[key]
        ws.append([
            key,
            stats["matches"],
            stats["wins"],
            stats["losses"],
            stats["win_rate"],
        ])

    ws.append([])
    ws.append(["Teammate", "Matches", "Wins", "Losses", "Win Rate"])
    for teammate in summary["teammates"]:
        ws.append(
            [
                teammate["full_name"],
                teammate["matches"],
                teammate["wins"],
                teammate["losses"],
                teammate["win_rate"],
            ]
        )

    ws.append([])
    ws.append(["Active Days"])
    for day in summary["calendar"]:
        ws.append([day])

    return _workbook_to_bytes(wb)


def export_all_matches(db: Session) -> bytes:
    matches: List[Match] = db.query(Match).order_by(Match.played_at.asc()).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "All Matches"
    ws.append([
        "Match ID",
        "Played At",
        "Session",
        "Type",
        "Score A",
        "Score B",
        "Winner",
        "Side A",
        "Side B",
        "Recorder",
    ])

    for match in matches:
        side_a = ", ".join(
            participant.user.full_name for participant in match.participants if participant.side == "A"
        )
        side_b = ", ".join(
            participant.user.full_name for participant in match.participants if participant.side == "B"
        )
        ws.append(
            [
                match.id,
                match.played_at.isoformat(sep=" "),
                match.session_number or "-",
                match.match_type.value,
                match.score_a,
                match.score_b,
                match.winner_side,
                side_a,
                side_b,
                match.recorded_by.full_name,
            ]
        )

    return _workbook_to_bytes(wb)
